"""ProductPulse: SKU data-quality + anomaly pipeline (Olist dataset).

Steps: load raw CSVs -> profile data quality -> clean -> build SKU table
-> add synthetic stock (seeded) -> detect anomalies (IQR + z-score)
-> save to SQLite -> write Excel summary report.

Run:  python pipeline.py
"""
from __future__ import annotations

import logging
import sqlite3
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
RAW = BASE / "data" / "raw"
DB_PATH = BASE / "data" / "productpulse.db"
REPORT_PATH = BASE / "reports" / "productpulse_summary.xlsx"

SEED = 42          # makes the synthetic stock + injected anomalies reproducible
IQR_K = 1.5        # Tukey fence multiplier
Z_THRESH = 3.0     # |z| above this is an outlier
MIN_GROUP = 20     # categories smaller than this fall back to global stats
INJECT_RATE = 0.01 # share of SKUs given each known (injected) anomaly type

log = logging.getLogger("productpulse")

FILES = {
    "products": "olist_products_dataset.csv",
    "items": "olist_order_items_dataset.csv",
    "trans": "product_category_name_translation.csv",
}


# --------------------------------------------------------------------------
# 1. Load
# --------------------------------------------------------------------------
def load_raw() -> dict[str, pd.DataFrame]:
    missing = [f for f in FILES.values() if not (RAW / f).exists()]
    if missing:
        raise FileNotFoundError(
            f"Missing in {RAW}: {missing}. Run `python download_data.py` first."
        )
    return {k: pd.read_csv(RAW / v) for k, v in FILES.items()}


# --------------------------------------------------------------------------
# 2. Data-quality profiling (on RAW data, before cleaning)
# --------------------------------------------------------------------------
def profile_quality(products: pd.DataFrame, items: pd.DataFrame) -> pd.DataFrame:
    issues: list[dict] = []

    def add(table, check, count, total, severity):
        if count > 0:
            issues.append(
                dict(
                    table=table,
                    check=check,
                    records=int(count),
                    pct=round(100 * count / total, 2) if total else 0.0,
                    severity=severity,
                )
            )

    n = len(products)
    for col in products.columns:
        add("products", f"missing {col}", products[col].isna().sum(), n,
            "high" if col == "product_id" else "medium")
    add("products", "duplicate product_id (SKU)", products.duplicated("product_id").sum(), n, "high")
    add("products", "weight <= 0", (products["product_weight_g"] <= 0).sum(), n, "medium")
    for dim in ["product_length_cm", "product_height_cm", "product_width_cm"]:
        add("products", f"{dim} <= 0", (products[dim] <= 0).sum(), n, "medium")

    m = len(items)
    add("order_items", "duplicate (order_id, order_item_id)",
        items.duplicated(["order_id", "order_item_id"]).sum(), m, "high")
    add("order_items", "price <= 0", (items["price"] <= 0).sum(), m, "high")
    add("order_items", "freight < 0", (items["freight_value"] < 0).sum(), m, "medium")
    add("order_items", "missing price", items["price"].isna().sum(), m, "high")
    add("order_items", "product_id not in products (orphan)",
        (~items["product_id"].isin(products["product_id"])).sum(), m, "high")

    no_sales = (~products["product_id"].isin(items["product_id"])).sum()
    add("products", "SKU with no sales", no_sales, n, "low")

    cols = ["table", "check", "records", "pct", "severity"]
    return pd.DataFrame(issues, columns=cols)


# --------------------------------------------------------------------------
# 3. Clean + build SKU table
# --------------------------------------------------------------------------
def clean_and_build(raw: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, dict]:
    products, items, trans = raw["products"], raw["items"], raw["trans"]
    stats = {"raw_skus": len(products), "raw_items": len(items)}

    p = products.drop_duplicates("product_id").copy()
    stats["duplicate_skus_removed"] = len(products) - len(p)
    p = p.rename(columns={
        "product_name_lenght": "name_length",
        "product_description_lenght": "description_length",
        "product_photos_qty": "photos_qty",
        "product_weight_g": "weight_g",
        "product_length_cm": "length_cm",
        "product_height_cm": "height_cm",
        "product_width_cm": "width_cm",
    })

    # category: translate PT -> EN, fall back to PT name, then 'unknown'
    p = p.merge(trans, on="product_category_name", how="left")
    p["category"] = (
        p["product_category_name_english"]
        .fillna(p["product_category_name"])
        .fillna("unknown")
    )
    p = p.drop(columns=["product_category_name_english"])
    p["missing_category"] = p["product_category_name"].isna()

    # physical attributes: flag, don't impute
    dims = ["weight_g", "length_cm", "height_cm", "width_cm"]
    p["invalid_dimensions"] = p[dims].isna().any(axis=1) | (p[dims] <= 0).any(axis=1)

    # order items: dedupe, drop invalid price, drop orphans
    it = items.drop_duplicates(["order_id", "order_item_id"])
    stats["duplicate_items_removed"] = len(items) - len(it)
    bad_price = it["price"].isna() | (it["price"] <= 0)
    stats["invalid_price_items_removed"] = int(bad_price.sum())
    it = it[~bad_price]
    orphan = ~it["product_id"].isin(p["product_id"])
    stats["orphan_items_removed"] = int(orphan.sum())
    it = it[~orphan]

    agg = it.groupby("product_id").agg(
        units_sold=("order_item_id", "size"),
        current_price=("price", "median"),   # median = robust "list price"
        min_price=("price", "min"),
        max_price=("price", "max"),
        price_std=("price", "std"),
        avg_freight=("freight_value", "mean"),
        n_sellers=("seller_id", "nunique"),
    ).reset_index()

    p = p.merge(agg, on="product_id", how="left")
    p["units_sold"] = p["units_sold"].fillna(0).astype(int)
    stats["clean_skus"] = len(p)
    stats["clean_items"] = len(it)
    return p, stats


# --------------------------------------------------------------------------
# 4. Synthetic stock + injected anomalies (ground truth for evaluation)
# --------------------------------------------------------------------------
def add_synthetic_stock_and_inject(df: pd.DataFrame, seed: int = SEED) -> pd.DataFrame:
    """Olist has no inventory data, so stock_level is SIMULATED from units sold.
    A small share of SKUs get known anomalies so detector recall can be measured."""
    rng = np.random.default_rng(seed)
    df = df.copy()
    n = len(df)

    base = df["units_sold"].clip(lower=1).to_numpy()
    df["stock_level"] = (base * rng.uniform(1.5, 6, n)).round().astype(int) + rng.poisson(2, n)
    df["injected_anomaly"] = ""

    sold = df.index[df["current_price"].notna()].to_numpy()
    k = max(1, int(INJECT_RATE * n))

    def pick(pool, exclude):
        pool = np.setdiff1d(pool, exclude)
        return rng.choice(pool, size=min(k, len(pool)), replace=False)

    used = np.array([], dtype=int)
    for label in ["price_spike", "price_drop", "negative_stock", "overstock"]:
        idx = pick(sold, used)
        used = np.concatenate([used, idx])
        if label == "price_spike":
            df.loc[idx, "current_price"] *= rng.uniform(15, 50, len(idx))
        elif label == "price_drop":
            df.loc[idx, "current_price"] *= rng.uniform(0.02, 0.07, len(idx))
        elif label == "negative_stock":
            df.loc[idx, "stock_level"] = -rng.integers(1, 50, len(idx))
        else:
            df.loc[idx, "stock_level"] = (
                df.loc[idx, "stock_level"].clip(lower=10) * rng.integers(15, 40, len(idx))
            )
        df.loc[idx, "injected_anomaly"] = label
    df["current_price"] = df["current_price"].round(2)
    return df


# --------------------------------------------------------------------------
# 5. Anomaly detection
# --------------------------------------------------------------------------
def outlier_flags(values: pd.Series, groups: pd.Series) -> pd.DataFrame:
    """IQR fences + z-score, computed within each group (category).
    Groups with fewer than MIN_GROUP valid values use global statistics."""
    d = pd.DataFrame({"v": values, "g": groups})
    gcount = d.groupby("g")["v"].transform("count")
    use_g = (gcount >= MIN_GROUP).to_numpy()

    g = d.groupby("g")["v"]
    q1 = np.where(use_g, g.transform(lambda s: s.quantile(0.25)), d["v"].quantile(0.25))
    q3 = np.where(use_g, g.transform(lambda s: s.quantile(0.75)), d["v"].quantile(0.75))
    mean = np.where(use_g, g.transform("mean"), d["v"].mean())
    std = np.where(use_g, g.transform("std"), d["v"].std())

    iqr = q3 - q1
    v = d["v"].to_numpy()
    with np.errstate(divide="ignore", invalid="ignore"):
        z = np.where(std > 0, (v - mean) / std, np.nan)
    return pd.DataFrame({
        "z": z,
        "iqr_high": v > q3 + IQR_K * iqr,
        "iqr_low": v < q1 - IQR_K * iqr,
        "z_high": z > Z_THRESH,
        "z_low": z < -Z_THRESH,
    }, index=d.index)


def detect_anomalies(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    # ---- pricing: outliers in log(price) within category
    valid_price = df["current_price"].notna() & (df["current_price"] > 0)
    pf = outlier_flags(np.log(df["current_price"].where(valid_price)), df["category"])
    df["price_z"] = pf["z"].round(3)
    df["price_iqr_flag"] = (pf["iqr_high"] | pf["iqr_low"]).astype(int)
    df["price_z_flag"] = (pf["z_high"] | pf["z_low"]).astype(int)
    price_high = pf["iqr_high"] | pf["z_high"]
    price_low = pf["iqr_low"] | pf["z_low"]

    # ---- stock: units of stock per unit sold ("cover"), upper tail only
    cover = df["stock_level"].clip(lower=0) / df["units_sold"].clip(lower=1)
    df["stock_cover"] = cover.round(3)
    sf = outlier_flags(np.log1p(cover), df["category"])
    df["stock_z"] = sf["z"].round(3)
    df["stock_iqr_flag"] = sf["iqr_high"].astype(int)
    df["stock_z_flag"] = sf["z_high"].astype(int)
    overstock = sf["iqr_high"] | sf["z_high"]

    # ---- business rules
    negative_stock = df["stock_level"] < 0
    top_sellers = df["units_sold"] >= df["units_sold"].quantile(0.95)
    stockout_risk = (df["stock_level"] == 0) & top_sellers

    flags = pd.DataFrame({
        "price_high": price_high,
        "price_low": price_low,
        "overstock": overstock,
        "negative_stock": negative_stock,
        "stockout_risk": stockout_risk,
    }).fillna(False).astype(bool)

    df["anomaly_type"] = flags.apply(
        lambda r: ",".join(c for c in flags.columns if r[c]), axis=1
    )
    df["is_anomaly"] = (df["anomaly_type"] != "").astype(int)
    df["price_anomaly"] = (price_high | price_low).fillna(False).astype(int)
    df["stock_anomaly"] = (overstock | negative_stock | stockout_risk).astype(int)

    # severity: 2 = both methods agree (or a hard rule broke), 1 = one method
    price_sev = df["price_iqr_flag"] + df["price_z_flag"]
    stock_sev = df["stock_iqr_flag"] + df["stock_z_flag"]
    rule_sev = (negative_stock | stockout_risk).astype(int) * 2
    df["severity"] = np.maximum.reduce([price_sev, stock_sev, rule_sev]).clip(max=2)
    return df


def evaluate_detection(df: pd.DataFrame) -> pd.DataFrame:
    """Recall of the detectors against the injected ground truth."""
    expected = {
        "price_spike": "price_high",
        "price_drop": "price_low",
        "negative_stock": "negative_stock",
        "overstock": "overstock",
    }
    rows = []
    for label, flag in expected.items():
        sub = df[df["injected_anomaly"] == label]
        caught = sub["anomaly_type"].str.contains(flag).sum()
        rows.append(dict(injected_type=label, injected=len(sub), detected=int(caught),
                         recall_pct=round(100 * caught / len(sub), 1) if len(sub) else None))
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------
# 6. SQLite
# --------------------------------------------------------------------------
VIEWS = """
DROP VIEW IF EXISTS v_category_summary;
CREATE VIEW v_category_summary AS
SELECT category,
       COUNT(*)                                   AS skus,
       SUM(units_sold)                            AS units_sold,
       ROUND(AVG(current_price), 2)               AS avg_price,
       SUM(is_anomaly)                            AS anomalies,
       ROUND(100.0 * SUM(is_anomaly) / COUNT(*), 2) AS anomaly_rate_pct
FROM sku_metrics
GROUP BY category;

DROP VIEW IF EXISTS v_anomalies;
CREATE VIEW v_anomalies AS
SELECT product_id, category, units_sold, current_price, stock_level, stock_cover,
       price_z, stock_z, anomaly_type, severity
FROM sku_metrics
WHERE is_anomaly = 1
ORDER BY severity DESC, ABS(COALESCE(price_z, 0)) DESC;
"""


def save_sqlite(df, dq, evaluation, stats) -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    run = pd.DataFrame([{**stats, "run_at": datetime.now().isoformat(timespec="seconds"),
                         "total_anomalies": int(df["is_anomaly"].sum())}])
    with sqlite3.connect(DB_PATH) as con:
        df.to_sql("sku_metrics", con, if_exists="replace", index=False)
        dq.to_sql("dq_issues", con, if_exists="replace", index=False)
        evaluation.to_sql("detection_eval", con, if_exists="replace", index=False)
        run.to_sql("run_log", con, if_exists="append", index=False)
        con.execute("CREATE INDEX IF NOT EXISTS idx_sku_product ON sku_metrics(product_id)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_sku_category ON sku_metrics(category)")
        con.executescript(VIEWS)


# --------------------------------------------------------------------------
# 7. Excel report (reads back from SQLite so SQL does the aggregation)
# --------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style(ws, widths: dict[str, int] | None = None) -> None:
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 45)


def write_excel_report(stats: dict) -> Path:
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(DB_PATH) as con:
        total = pd.read_sql("SELECT COUNT(*) n, SUM(is_anomaly) a, SUM(price_anomaly) p, "
                            "SUM(stock_anomaly) s, SUM(missing_category) mc, "
                            "SUM(invalid_dimensions) idim FROM sku_metrics", con).iloc[0]
        dq = pd.read_sql("SELECT * FROM dq_issues ORDER BY "
                         "CASE severity WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END, records DESC", con)
        cats = pd.read_sql("SELECT * FROM v_category_summary ORDER BY anomalies DESC", con)
        anomalies = pd.read_sql("SELECT * FROM v_anomalies LIMIT 1000", con)
        evaluation = pd.read_sql("SELECT * FROM detection_eval", con)

    n = int(total["n"])
    summary = pd.DataFrame([
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Raw SKUs", stats["raw_skus"]),
        ("Duplicate SKUs removed", stats["duplicate_skus_removed"]),
        ("Clean SKUs", n),
        ("Raw order items", stats["raw_items"]),
        ("Duplicate order items removed", stats["duplicate_items_removed"]),
        ("Invalid-price items removed", stats["invalid_price_items_removed"]),
        ("SKUs missing category", int(total["mc"])),
        ("SKUs with invalid/missing dimensions", int(total["idim"])),
        ("Total anomalies flagged", int(total["a"])),
        ("Anomaly rate", f"{100 * total['a'] / n:.2f}%"),
        ("Pricing anomalies", int(total["p"])),
        ("Stock anomalies", int(total["s"])),
        ("Method", f"IQR (k={IQR_K}) + z-score (|z|>{Z_THRESH}), per category on log scale"),
        ("Note", "stock_level is synthetic (seeded) - Olist has no inventory data"),
    ], columns=["Metric", "Value"])

    with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        dq.to_excel(xw, sheet_name="Data Quality", index=False)
        cats.to_excel(xw, sheet_name="Category Summary", index=False)
        anomalies.to_excel(xw, sheet_name="Top Anomalies", index=False)
        evaluation.to_excel(xw, sheet_name="Detection Eval", index=False)

        for ws in xw.book.worksheets:
            _style(ws)
        xw.book["Summary"].column_dimensions["A"].width = 38
        xw.book["Summary"].column_dimensions["B"].width = 70

        # severity highlighting
        ws = xw.book["Data Quality"]
        colours = {"high": "F8CBAD", "medium": "FFE699", "low": "E2EFDA"}
        for row in ws.iter_rows(min_row=2):
            fill = colours.get(row[4].value)
            if fill:
                row[4].fill = PatternFill("solid", fgColor=fill)

        # chart: top 10 categories by anomaly count
        ws = xw.book["Category Summary"]
        top = min(11, ws.max_row)
        chart = BarChart()
        chart.type, chart.title = "bar", "Top 10 categories by anomalies"
        chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=top), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=top))
        chart.height, chart.width = 9, 18
        ws.add_chart(chart, "H2")

    return REPORT_PATH


# --------------------------------------------------------------------------
def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(message)s", datefmt="%H:%M:%S")
    raw = load_raw()
    log.info("Loaded %d products, %d order items", len(raw["products"]), len(raw["items"]))

    dq = profile_quality(raw["products"], raw["items"])
    log.info("Data-quality checks with findings: %d", len(dq))

    df, stats = clean_and_build(raw)
    df = add_synthetic_stock_and_inject(df)
    df = detect_anomalies(df)
    evaluation = evaluate_detection(df)

    save_sqlite(df, dq, evaluation, stats)
    report = write_excel_report(stats)

    log.info("Anomalies flagged: %d of %d SKUs (%.2f%%)",
             df["is_anomaly"].sum(), len(df), 100 * df["is_anomaly"].mean())
    log.info("Detection recall on injected anomalies:\n%s", evaluation.to_string(index=False))
    log.info("SQLite: %s", DB_PATH)
    log.info("Excel : %s", report)


if __name__ == "__main__":
    main()
